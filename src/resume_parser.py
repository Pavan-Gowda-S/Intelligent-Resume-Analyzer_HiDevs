"""
Resume parsing engine.

Supports PDF (via PyMuPDF / fitz), DOCX (via python-docx), and TXT files.
All parsing happens in-memory from bytes; no temporary files are written.
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Any

from src.config import MAX_FILE_SIZE_BYTES, SUPPORTED_EXTENSIONS
from src.models import Education, ParsedResume, WorkExperience
from src.skill_extractor import extract_skills, group_skills_by_category
from src.text_processor import (
    estimate_total_experience_months,
    extract_candidate_name,
    extract_certifications,
    extract_education_entries,
    extract_email,
    extract_experience_entries,
    extract_phone,
    segment_sections,
)
from src.utils import normalize_text
from src.validators import validate_file_extension, validate_file_size

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Low-level text extraction
# ---------------------------------------------------------------------------

def _extract_text_pdf(data: bytes) -> tuple[str, str, list[str], bool]:
    """
    Extract text from PDF bytes using PyMuPDF.

    Returns (text, method_description, warnings, is_scanned).
    """
    warnings: list[str] = []
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(stream=data, filetype="pdf")
        pages_text: list[str] = []
        for page in doc:
            pages_text.append(page.get_text())  # type: ignore[arg-type]
        doc.close()

        full_text = "\n".join(pages_text)
        stripped = full_text.strip()

        if not stripped:
            warnings.append(
                "This PDF appears to be a scanned image or contains no "
                "extractable text. OCR is not available in this version. "
                "Please provide a text-based PDF or a DOCX/TXT file."
            )
            return "", "pymupdf (no text — possible scan)", warnings, True

        return full_text, "pymupdf", warnings, False

    except ImportError:
        warnings.append(
            "PyMuPDF is not installed. Falling back to pdfminer.six for PDF parsing."
        )
        return _extract_text_pdf_fallback(data)
    except Exception as exc:
        msg = str(exc)
        if "encrypted" in msg.lower() or "password" in msg.lower():
            warnings.append(
                "This PDF is encrypted/password-protected and cannot be parsed."
            )
        else:
            logger.warning("PDF parsing error: %s", msg)
            warnings.append(f"PDF parsing encountered an error: {msg}")
        return "", "pymupdf (error)", warnings, False


def _extract_text_pdf_fallback(data: bytes) -> tuple[str, str, list[str], bool]:
    """Fallback PDF extraction using pdfminer.six."""
    from pdfminer.high_level import extract_text as pdfminer_extract

    warnings: list[str] = []
    try:
        text = pdfminer_extract(io.BytesIO(data))
        if not text or not text.strip():
            warnings.append("No text could be extracted from this PDF.")
            return "", "pdfminer (no text)", warnings, True
        return text, "pdfminer", warnings, False
    except Exception as exc:
        warnings.append(f"pdfminer parsing error: {exc}")
        return "", "pdfminer (error)", warnings, False


def _extract_text_docx(data: bytes) -> tuple[str, str, list[str]]:
    """Extract text from DOCX bytes using python-docx."""
    from docx import Document

    warnings: list[str] = []
    try:
        doc = Document(io.BytesIO(data))
        parts: list[str] = [para.text for para in doc.paragraphs]

        # Also extract table cell text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)

        full_text = "\n".join(parts)
        if not full_text.strip():
            warnings.append("The DOCX file appears to contain no readable text.")
        return full_text, "python-docx", warnings

    except Exception as exc:
        logger.warning("DOCX parsing error: %s", exc)
        warnings.append(f"DOCX parsing error: {exc}")
        return "", "python-docx (error)", warnings


def _extract_text_txt(data: bytes) -> tuple[str, str, list[str]]:
    """Extract text from raw bytes with encoding detection fallback."""
    warnings: list[str] = []
    for enc in ("utf-8", "utf-16", "latin-1", "cp1252"):
        try:
            text = data.decode(enc)
            if text.strip():
                return text, f"plain-text ({enc})", warnings
        except (UnicodeDecodeError, ValueError):
            continue
    warnings.append("Could not decode file with any supported encoding.")
    return "", "plain-text (decode error)", warnings


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def parse_resume(file_bytes: bytes, filename: str) -> ParsedResume:
    """
    Parse a resume from raw bytes.

    This is the main entry point.  It validates inputs, extracts raw text,
    then delegates to the text-processing pipeline.

    Parameters
    ----------
    file_bytes : bytes
        Raw file content.
    filename : str
        Original filename, used only for extension detection and record-keeping.

    Returns
    -------
    ParsedResume
        A fully populated (or partially populated with warnings) result object.
        The ``raw_text`` field is set; never call this with untrusted filenames
        for path-related operations.
    """
    resume = ParsedResume(filename=filename)

    # --- File-level validation -------------------------------------------
    ext = Path(filename).suffix.lower()
    resume.file_format = ext.lstrip(".")

    ok, msg = validate_file_size(file_bytes, filename)
    if not ok:
        resume.parse_success = False
        resume.warnings.append(msg)
        return resume

    ok, msg = validate_file_extension(filename)
    if not ok:
        resume.parse_success = False
        resume.warnings.append(msg)
        return resume

    # --- Text extraction -------------------------------------------------
    raw_text = ""
    method = ""

    if ext == ".pdf":
        raw_text, method, pdf_warnings, is_scan = _extract_text_pdf(file_bytes)
        resume.warnings.extend(pdf_warnings)
        resume.is_scanned_pdf = is_scan
        if is_scan:
            resume.parse_success = False
            resume.text_extraction_method = method
            resume.raw_text = ""
            return resume

    elif ext == ".docx":
        raw_text, method, docx_warnings = _extract_text_docx(file_bytes)
        resume.warnings.extend(docx_warnings)

    elif ext == ".txt":
        raw_text, method, txt_warnings = _extract_text_txt(file_bytes)
        resume.warnings.extend(txt_warnings)

    resume.text_extraction_method = method

    if not raw_text.strip():
        resume.parse_success = False
        if not resume.warnings:
            resume.warnings.append("No text could be extracted from this file.")
        return resume

    # --- Text normalisation & field extraction ---------------------------
    resume.raw_text = raw_text
    text_norm = normalize_text(raw_text)

    # Contact info
    resume.email = extract_email(text_norm)
    resume.phone = extract_phone(text_norm)

    # Candidate name
    name, confidence = extract_candidate_name(raw_text, filename)
    resume.candidate_name = name
    resume.name_confidence = confidence
    if confidence == "unknown":
        resume.warnings.append(
            "Candidate name could not be reliably detected. "
            "Please verify manually."
        )

    # Section segmentation
    sections = segment_sections(raw_text)

    # Skills (extracted from full text to maximise coverage)
    resume.skills = extract_skills(text_norm)
    resume.skills_by_category = group_skills_by_category(resume.skills)

    # Education
    edu_section = sections.get("education", "") + "\n" + sections.get("other", "")
    raw_edu_entries = extract_education_entries(edu_section)
    resume.education = [Education.from_dict(e) for e in raw_edu_entries]

    # Certifications
    cert_section = sections.get("certifications", "") + "\n" + text_norm
    resume.certifications = extract_certifications(cert_section)

    # Work experience
    exp_section = sections.get("experience", sections.get("other", text_norm))
    raw_exp_entries = extract_experience_entries(exp_section)
    resume.work_experience = [WorkExperience.from_dict(e) for e in raw_exp_entries]

    # Experience estimation
    resume.total_experience_months = estimate_total_experience_months(raw_exp_entries)
    if resume.total_experience_months is None:
        resume.warnings.append(
            "Employment dates are unclear or missing; total experience could not "
            "be estimated reliably. This will affect the experience component of "
            "the match score."
        )

    return resume


def parse_resume_from_path(path: str | Path) -> ParsedResume:
    """Convenience wrapper for loading from a filesystem path."""
    p = Path(path)
    with open(p, "rb") as fh:
        data = fh.read()
    return parse_resume(data, p.name)
